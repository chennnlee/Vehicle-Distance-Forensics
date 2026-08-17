#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""大專學生研究計畫成果報告草稿產生器(python-docx)。

2026-08-17 更新:補上 7/12 之後的成果——公開資料集驗證(comma2k19)、
KITTI 取樣率邊界、基頻選擇改為跨時間聯合決策,以及一次誤差歸因的公開更正。
同時補上申請人／系所／指導教授的佔位列(原版沒有這些必填欄位)。
7/12 版本保留為 build_report_20260712.py.bak,以便對照。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# 圖片素材(17 MB)不進版控——它們是從 data/output/ 的渲染輸出與影格裁出來的。
# 預設路徑沿用既有位置,可用 REPORT_ASSETS 覆寫。產生方式見本目錄 README.md。
A = os.environ.get('REPORT_ASSETS', '/home/s11244/tmp/pptx_build/assets')
OUT_DIR = '/home/s11244/code/114/Vehicle-Distance-Forensics/data/output/report_draft'
os.makedirs(OUT_DIR, exist_ok=True)

doc = Document()

# ---- page & base style (A4, 2.5cm margins) ----
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.5)

def set_font(style_or_run, latin='Times New Roman', east='標楷體', size=None, bold=None, color=None):
    f = style_or_run.font
    f.name = latin
    r = style_or_run.element.rPr if hasattr(style_or_run, 'element') else None
    # ensure eastAsia
    rpr = f.element.rPr if hasattr(f, 'element') else None
    try:
        f._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    except Exception:
        pass
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color: f.color.rgb = RGBColor(*color)

st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '標楷體')
st.font.size = Pt(12)
pf = st.paragraph_format
pf.line_spacing = 1.5
pf.space_after = Pt(6)

def para(text='', size=12, bold=False, align=None, east='標楷體', latin='Times New Roman',
         indent=None, space_after=6, color=None, line=1.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(indent)
    r = p.add_run(text)
    r.font.name = latin
    r.element.rPr.rFonts.set(qn('w:eastAsia'), east)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p

def h1(text):
    return para(text, size=16, bold=True, east='微軟正黑體', latin='Arial', space_after=8)

def h2(text):
    return para(text, size=13.5, bold=True, east='微軟正黑體', latin='Arial', space_after=6)

def body(text):
    return para(text, size=12, indent=0.85)

def fig(path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = para(caption, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    for r in c.runs:
        r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return p

def add_table(headers, rows, widths_cm, caption=None, font_size=10.5):
    if caption:
        cp = para(caption, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htxt in enumerate(headers):
        cell = t.cell(0, j)
        cell.width = Cm(widths_cm[j])
        cp2 = cell.paragraphs[0]
        r = cp2.add_run(htxt)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.name = 'Arial'
        r.element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
        sh = cell._tc.get_or_add_tcPr()
        el = sh.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'EFF2F7'})
        sh.append(el)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i + 1, j)
            cell.width = Cm(widths_cm[j])
            r = cell.paragraphs[0].add_run(str(val))
            r.font.size = Pt(font_size)
            r.font.name = 'Times New Roman'
            r.element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')
    para('', size=6, space_after=6)
    return t

# ================= 封面/標題 =================
para('國家科學及技術委員會 補助大專學生研究計畫', size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
     east='微軟正黑體', latin='Arial', space_after=2)
para('研究成果報告(草稿)', size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
     east='微軟正黑體', latin='Arial', space_after=14)
para('AI 單鏡頭影像之交通事故車輛速度與距離量測', size=20, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, east='微軟正黑體', latin='Arial', space_after=4)
para('Vehicle-Distance-Forensics:單幀 3D 幾何 × 台灣法定標線尺度錨 × 誠實不確定度',
     size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20, color=(0x55, 0x55, 0x55))
para('申請人:＿＿＿＿＿＿　　系所／年級:＿＿＿＿＿＿　　指導教授:＿＿＿＿＿＿',
     size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para('執行期間:＿＿＿＿年＿＿月 ～ ＿＿＿＿年＿＿月　(草稿版 2026-08-17)', size=11,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30, color=(0x55, 0x55, 0x55))

# ================= 摘要 =================
h1('摘要')
body('交通事故鑑定實務中,警方與鑑識單位須自路口監視器(CCTV)與行車紀錄器影像中'
     '還原車輛速度與距離,作為責任釐清與法庭審理之科學依據。然而此類影像均為單鏡頭'
     '系統,缺乏直接測距能力;現行人工畫格法費時且高度依賴專家經驗。本計畫建立一套'
     '以單幀 3D 重建為幾何骨架、以台灣法定標線為尺度錨的自動化量測系統,涵蓋固定式'
     ' CCTV 與移動式行車紀錄器兩條管線,輸出附 ±95% 信賴區間之車速與距離,並以'
     '「誠實不確定度」為核心設計原則——量測不可信時自動拒發,絕不硬給數字。')
body('驗證成果:CCTV 管線與人工交比法同窗比對誤差 +0.5%(前鎮 1130221 案);'
     '行車紀錄器自車速對 GPS 之 MAE 為 1~2.4 km/h;20251029 案例包 8 件汽車行車'
     '紀錄器鑑定案批量驗證(28 個量測段)對人工畫格法 MAE 1.93 km/h,剔除經 GPS'
     ' 仲裁判定人工真值本身可疑之 5 段後,MAE 1.37 km/h,±3 km/h 內達 24/28 段。'
     '過程中並系統性揭露證物影片之三類「時基陷阱」(容器幀率謊報、變動幀率塞幀、'
     '剪輯跳接),提出以畫面內 OSD 時鐘重建時間軸之稽核流程,對鑑識實務具直接參考價值。')
body('本階段(2026 年 8 月)另將驗證延伸至附論文與儀器真值之公開資料集,'
     '使精度可被外部檢驗:於 comma2k19(Schafer 等,2018;加州通勤,同時提供 CAN 車速'
     '與 GNSS/INS 全域位姿)之 24 段 / 24 分鐘 / 28,703 幀上,自車速對 GNSS/INS 位姿'
     '真值之 MAE 為 2.79 km/h(高速公路、日夜皆含;僅日間 1.78,中位相對誤差 0.98%),'
     '且全程未使用任何真值校正。過程中另發現行車紀錄器自帶 GPS 並非金標準——'
     '實測偏差達 ±3 km/h 級且逐機種不同,本系統反而更接近人工畫格法真值。')
para('關鍵詞:交通事故鑑識、單眼深度估計、尺度校正、法定標線、車速量測、不確定度量化、公開資料集驗證',
     size=12, indent=0.85, space_after=12)

doc.add_page_break()

# ================= 一、研究動機 =================
h1('一、研究動機與研究問題')
body('依道安統計,全臺交通事故案件量逐年攀升,肇事責任釐清高度依賴影像證據。'
     '碰撞前的行駛速度與距離變化是判定路權與過失比例的關鍵參數,但現行程序仰賴'
     '人工逐幀調閱影像(畫格法),不僅耗時費力,且受影像畫質、光影與主觀認定影響。'
     '固定式 CCTV 之高視角俯瞰帶來透視變形;行車紀錄器則因相機本身移動,使以地面'
     '固定參考物為準之傳統二維量測法失效。')
body('本計畫的核心研究問題為:如何在「相機內參未知、尺度歧義、時間軸不可信」'
     '三重限制下,自單鏡頭影像產出具司法信度的速度與距離量測?對應三項具體目標:'
     '(1) 建立單鏡頭深度量測之標準化驗證流程,以法定物理錨點量化誤差並定義失效邊界;'
     '(2) 研發具物理幾何約束之自適應尺度校正機制;(3) 強化移動式行車紀錄器之鑑識'
     '應用,處理自車與目標車之相對運動。')
body('本階段新增第四項目標:將驗證自「自有案例」延伸到「有論文、有儀器真值的公開'
     '資料集」,使量測精度可被外部檢驗,並可與文獻直接比較。此舉同時暴露了一條'
     '原本被素材條件掩蓋的適用條件(見 4.7 節),對方法本身的可靠性陳述有實質影響。')

# ================= 二、文獻回顧 =================
h1('二、文獻回顧與方法演進')
body('傳統幾何投影法以消失點或單應性矩陣將影像座標映射回路面座標,依賴法定標線'
     '作為比例尺,但在視角偏斜與遠場情境誤差急遽放大。三維模型比對法需車輛 3D 模型庫'
     '與高昂建模成本,泛化性受限。深度學習單眼深度估計(MDE)如 Metric3D v2、'
     'Depth Anything v2、UniDepth v2 提供了泛化的深度先驗,但本計畫初期誤差量化實驗'
     '發現:此類模型在真實交通場景存在尺度歧義、幀間深度跳動與領域偏差,'
     '原始輸出不足以直接作為鑑定數據。')
body('據此,本計畫方法演進為:以單幀 3D 高斯重建(SHARP)取得幾何形狀,'
     '「形狀信任模型、尺度信任法定標線」——模型只負責幾何骨架,絕對尺度一律由'
     '台灣法定標線(車道虛線 4m/6m、高速公路 4m/8m、路口導引線 0.5m/0.5m)反推,'
     '並以跨鏈互驗之變異係數(CV)將校正不確定度傳遞至最終速度的信賴區間。'
     '此設計繞開了相機內參未知與模型尺度漂移兩大瓶頸,亦即申請書中'
     '「AI 基礎模型 × 物理幾何約束」雙軌校正構想之落實。')
body('在速度量測的公開評比方面,Fernández Llorca 等(IET ITS, 2021)之綜述指出,'
     '附速度真值的公開資料集極少,主要為 BrnoCompSpeed(Sochor 等, IEEE T-ITS 2019;'
     '20,865 台車、雷射光閘真值、附相機標定)與 UTFPR(Luvizon 等, 2017),兩者皆為'
     '固定式相機。行車紀錄器自車速方面,comma2k19 提供 CAN 車速與 GNSS/INS 位姿,'
     '且路型為連續虛線之高速公路,最貼近本計畫之碼表法。本計畫據此規劃:管線 A 對應'
     'BrnoCompSpeed,管線 B 對應 comma2k19;KITTI 雖有精確速度真值,但其 10 Hz 取樣率'
     '在高速段不足(見 4.8 節之取樣率條件)。')

# ================= 三、研究方法 =================
h1('三、研究方法及步驟')
h2('3.1 幾何骨架:單幀 3D 重建與地面平面')
body('整段影片僅取一張參考幀執行 SHARP 單目 3D 重建(CPU 約 90 秒,全片僅需一次),'
     '於點雲中以 RANSAC 擬合地面平面 y = ax + bz + c;此後任意像素之路面座標皆由'
     '「相機射線 × 平面求交 × 尺度」即時求得,逐幀量測不再依賴深度模型,'
     '確保時序穩定性。圖 1 為 dc007 案之實際重建輸出。')
fig(f'{A}/sharp_pointcloud.png', 12.5,
    '圖 1:SHARP 單幀點雲(可辨識路面虛線、紅色目標車、門架)與 RANSAC 地面平面(琥珀色)')

h2('3.2 尺度錨:台灣法定標線')
body('SHARP 採固定 FOV 假設,尺度天生不可信(實測各裝置 scale 介於 0.88~2.9)。'
     '本系統以法定標線反推尺度:先以 gap/dash 比例做免尺度的標線規格識別,再以'
     '多鏈剝離讓每條虛線各自成為一把獨立的尺,跨鏈互驗得到尺度變異 CV,直接進入'
     '速度不確定度。實測 4m 車道線跨鏈一致性約 5%,優於 0.5m 導引線(10~22%),'
     '故能用 4m 線就用 4m 線。行車紀錄器因貼地視角使虛線 3D 取樣失效,改以車道寬'
     '(3.5m)、同鏡頭相機高度轉移或 GPS 作錨。')
fig(f'{A}/64000C000013_lane_dash_calibration.png', 12.5,
    '圖 2:kh013 路口 CCTV 之標線校正——SHARP 原始 3D 長度 × scale 2.680 ≈ 法定 4m(CV 7.4%)')

h2('3.3 管線 A:固定式 CCTV 測速')
body('流程:(1) OSD 時鐘秒跳定時——NVR 串流謊報幀率是常態,真實時間軸以畫面'
     '時鐘重建;(2) YOLOv8m-seg 以車輛遮罩最低 12% 取接地點,映射為路面公尺座標;'
     '(3) 匈牙利關聯佐以方向閘、bbox 尺寸閘與物理閘門(速度誤差項 + 4 m/s² 加速度上限)'
     '串接軌跡,並具同幀重複框抑制、類別群組關聯、斷軌事後縫合與遠場防護鏈'
     '(v3 追蹤器);(4) 等速擬合輸出速度與 ±95% 信賴區間,非等速標示區間平均,'
     '遠場(靈敏度 >0.3 m/px)自動拒發。')
fig(f'{A}/kh013_t25.jpg', 13.5,
    '圖 3:管線 A 實際輸出(kh013 夜間路口)——校正標線、車輛軌跡、速度標籤與 HUD 全部燒錄於影片')

h2('3.4 管線 B:行車紀錄器(移動相機)')
body('自車速唯一信任來源為「虛線週期碼表」:法定虛線以固定週期流過畫面固定列,'
     '對該列亮度訊號做滑動自相關取週期 T,自車速 = 週期長 ÷ T。此法先天免疫於尺度、'
     '鏡頭畸變與車身震動。訊號品質由十道防護鏈把關(停等假鎖、窗內週期數、雙窗一致、'
     '峰突出度、跨列相位、八度錯誤、線位穩定、子結構半速、分裂鎖、光流禁用),'
     '每一道閘門都對應一支實際失效影片;全部擋完仍無讀值則誠實標 no-lock。'
     '他車量測:接地點 → 平面 → 跟車距離;絕對速度 = 自車速 + 距離變化率。')
body('取樣點選定方法(dc003 案教訓):先將「畫面某列亮度」沿時間堆疊為線位×時間'
     '條紋圖,白色點狀縱列即虛線脈衝列,取樣點定於駐留平台上,不憑單幀肉眼猜測;'
     '同一張圖亦可辨識虛轉實線、雙白線與減速標線干擾帶,是拒發窗的可稽核證據(圖 4)。')
body('本階段另完成取樣點之自動挑選工具,使「整個資料集跑一遍」成為可能——'
     '公開資料集評估不可能逐段人工瞄點。與人工挑點於同一支影片上比對,'
     '逐幀速度差之中位數為 0.39 km/h,實質等價。')
fig(f'{A}/strip_c0135.png', 9.5,
    '圖 4:案 01-35 之線位×時間條紋圖——f0–f130 駐留平台清楚,其後車輛道內右漂')
fig(f'{A}/dc006_t45.jpg', 13.5,
    '圖 5:管線 B 實際輸出(dc006 國道)——自車視覺碼表 90.0 vs 畫面 GPS 92 km/h;前車逐幀距離與絕對速度')

h2('3.5 證物影片時基稽核(重要流程)')
body('速度 = 距離 ÷ 時間,而證物影片的「時間」不可信。本計畫在批量驗證中系統性'
     '遭遇三類時基陷阱,並建立稽核流程:先以 ffprobe 檢視容器幀率與 pts 分布,'
     '再與畫面 OSD 時鐘比對(人工報表之 FTS 慣例上即為「總幀數 ÷ OSD 時鐘跨度」),'
     '必要時對秒位數字逐幀分群偵測跳變,重建真實時間軸後才進行任何速度計算。'
     '詳細案例見 4.3 節。')

h2('3.6 基頻選擇:從逐幀獨立改為跨時間聯合決策')
body('亮度自相關無法分辨「真週期」與其整數倍或整數分之一,這是本方法的固有歧義。'
     '舊做法在單一時間窗內以諧波梳狀評分挑選基頻;本階段在公開資料集上發現該做法'
     '於夜間會系統性失效——路面逆反射標記在車燈照射下遠比漆線明亮,使「半週期」'
     '成為自相關上真實存在的峰。我們量測了五個脈衝形狀特徵(時間寬度、峰銳度、面積、'
     '振幅、縱向跨列高度)試圖分辨兩種脈衝,全部失敗:誤鎖窗之交替程度並不高於正確窗'
     '(相鄰對一致率 0.53~0.72,隨機為 0.50),而物理上最應成立的縱向跨列高度,'
     '兩組中位數僅差 6.0 對 7.0 列——夜間長曝光將點狀標記拖成條紋,正好抹除該特徵。'
     '結論是:在單一時間窗之內無法選擇基頻。')
body('解法為改變決策的單位:車輛不可能在一幀之內將速度變成兩倍。新做法對每個取樣點'
     '沿時間執行 Viterbi 動態規劃,發射項沿用原本有效之諧波梳狀評分,轉移項對隱含的'
     '速度變化收費;候選僅取該幀自相關真正存在之峰,因此真實加減速仍可自由跟隨,'
     '變得昂貴的只有「跳一個八度再跳回來」。實作上有一關鍵細節:轉移項必須以速度的'
     '對數變化計費,不可使用絕對速度差——低一個八度的路徑其速度變化亦低一個八度,'
     '永遠看似更平滑,罰絕對值等於獎勵減半(第一版即因此打壞既有案例,'
     '某案覆蓋率自 57.8% 降至 27.1%)。改為對數尺度後,真值路徑與減半路徑代價相同,'
     '基頻回歸由證據決定。成效與既有案例之迴歸驗證見 4.6 與 4.9 節。')

# ================= 四、研究結果 =================
doc.add_page_break()
h1('四、研究結果與討論')

h2('4.1 管線 A 驗證:對人工交比法')
body('前鎮 1130221 案(白天路口 CCTV,見圖 6):與人工交比法同時間窗比對,'
     '本系統 48.6 km/h、人工 48.37 km/h,誤差 +0.5%;幀對幀互查顯示人工所選幀號'
     '恰為系統幀號 −1,輪胎接地座標互相吻合。海盛提供之 5 件 CCTV 案中 3 案吻合'
     '(45.5 / 48.4 / 44.6 km/h),2 案屬系統自動拒發之失效邊界(近地平線遠場、'
     '螢幕翻拍摩爾紋);其中 1110822 案兩條軌跡靈敏度 0.83/0.88 m/px,'
     '系統自動標記 far_field,驗證了遠場防護之有效性。')
fig(f'{A}/qz_t30.jpg', 13.0,
    '圖 6:前鎮 1130221 案(白天路口)——0.5m 導引線錨、軌跡與即時速度;遠場車輛誠實標 far-range')

h2('4.2 管線 B 驗證:純視覺自車速對 GPS 與人工畫格法')
body('多段行車紀錄器影片以畫面燒錄之 GPS 車速為對照:002 快速道路 11 檢查點 MAE'
     ' 1.2 km/h;dc006(國道一號,60 秒)9 有效檢查點 MAE ≈ 2.1;dc007(台86)'
     'MAE ≈ 2.2;dc003(超廣角、自車三度換道)鎖定段多在 ±4 km/h 內,碼表覆蓋率'
     ' 52% 誠實呈現。急煞與換道期間依設計拒發(no-lock),不硬給數字。')
body('有人工畫格法真值者一律以人工為準:hs005(國道巡航 18 秒)同窗比對 12 段,'
     'MAE 0.69 km/h、bias +0.01、12/12 落在 ±3 km/h 內;hs006s 為 MAE 1.30。'
     '此外四方交叉(GPS 座標差分 93.7、GPS 速度欄 94.17、人工 94.42、本系統 94.83)'
     '散布於 1.1 km/h 之內。')
body('本階段的一項重要發現是:行車紀錄器自帶 GPS 並非金標準。取兩支同時具備 GPS OSD'
     '與人工畫格法真值之影片做三方比對,本系統均比該機自身之 GPS 更接近人工真值,'
     '且 GPS 偏差逐機種不同——一機無偏但散(MAE 2.09),另一機系統性低報 2.89 km/h。'
     '推論:過去所有「對照 GPS」之驗證數字中有一部分實為裝置誤差,本系統實際精度'
     '應優於已記錄值。凡僅有 OSD 可對照者,均須註明對照對象本身有 ±3 km/h 級不確定度。')
fig(f'{A}/ego_vs_gps_validation.png', 13.5,
    '圖 7:純視覺碼表(藍)對 GPS(綠點)——高速巡航緊貼,無法鎖定處誠實留空')

h2('4.3 批量驗證:20251029 汽車 8 案 28 段')
body('本計畫對 20251029 案例包全部 8 件汽車行車紀錄器鑑定案完成批量驗證,'
     '共 28 個量測段(表 1、圖 8):全體 MAE 1.93 km/h;經畫面 GPS 仲裁,'
     '5 個大偏差段之誤差來源在人工端(±1 幀量化、時基誤用),剔除後 23 段 MAE'
     ' 1.37 km/h;±3 km/h 內 24/28 段。夜間案 33-23 全程 100% 鎖定(conf 0.65),'
     '顯示碼表於夜間路燈環境依然可用。')
add_table(
    ['案', '條件', '段數', '對人工差值(km/h)', '備註'],
    [
        ['78-28', '夜 576p', '2', '+2.8 / −1.3', '2026-07-10 完成'],
        ['78-29', '夜 1296p', '1(+1 no-lock)', '+2.1', '片尾窗截斷誠實 no-lock'],
        ['01-35', '晝雨後 25fps', '5', '+2.3 −3.8 −4.5 +0.7 −1.7', 'GPS 平滑減速 77→70;人工 75/75 為幀量化'],
        ['43-16', '黃昏 30.5fps', '4', '+0.6 −2.5 −1.7 −3.0', '全程加速;碼表 chirp 平滑(已知邊界)'],
        ['10-05', '黃昏 29.85fps', '3', '−4.0 −2.5 −0.6', 'GPS 仲裁:seg1/2 本系統較準'],
        ['33-23', '夜 28.75fps', '4', '−0.7 −1.6 +0.4 −8.0', 'seg4 人工 86.25 vs GPS 80–81:人工偏高'],
        ['14-16', '晝 VFR', '3', '+1.8 +2.0 −0.4', '對 pts 修正真值;人工發布值高 +30%'],
        ['44-07', '清晨 30fps', '6', '+2.5 +1.2 −0.1 +1.2 −0.3 −0.1', '對 ×2.899 修正真值;剪輯掉 60 秒'],
    ],
    [1.7, 2.6, 2.4, 4.6, 4.6],
    caption='表 1:8 件汽車行車紀錄器案批量驗證彙總(對人工畫格法)')
fig(f'{A}/error_distribution.png', 15.0,
    '圖 8:28 段誤差分布——紅圈為 GPS 仲裁判定人工真值可疑之段;右側直方圖 MAE 全體 1.93 / 可信 1.37')

h2('4.4 發現:證物影片的三類時基陷阱')
body('批量驗證同時揭露了對鑑識實務具普遍意義的發現——證物影片的時間軸不可信,'
     '三類陷阱在本批案例全部實際出現(表 2)。原則:人工畫格法之虛線計數(距離)'
     '通常正確,誤差幾乎都來自時基;與 CCTV 管線「NVR 謊報幀率 → OSD 秒跳定時」'
     '為同一原則。案 44-07 之逐幀時鐘稽核(圖 9)顯示影片於 f810/f811 間剪掉整整'
     ' 60 秒(07:25:36 → 07:26:36,秒位連續、分位才露餡),人工將 92 秒攤到 952 幀,'
     '速度被低估 2.9 倍;依 ×2.899 修正後,人工距離、OSD 車速與本系統三方吻合'
     '(六段 MAE 0.88 km/h)。')
add_table(
    ['陷阱', '案例', '現象', '處理'],
    [
        ['容器幀率謊報', '33-23', '容器 29.97fps 且 pts 等距,OSD 時鐘 8 秒 230 幀 → 真時基 28.75fps(差 4.2%)', '以 OSD 時鐘重建幀率'],
        ['VFR 爆發塞幀', '14-16', '名義 13fps,實為 10fps 主幀+爆發塞幀;人工均勻假設使發布值高 +30%', '以容器 pts 重取樣為均勻格點(GPS 互證)'],
        ['剪輯跳接', '44-07', 'f810/811 間剪掉 60 秒;人工把 OSD 總跨度攤到全幀數 → 速度 ÷2.9', '秒位逐幀分群偵測跳變,切出連續段(30fps)'],
    ],
    [2.6, 1.6, 6.6, 5.1],
    caption='表 2:證物影片時基三陷阱與處理')
fig(f'{A}/c4407_cut.png', 5.5,
    '圖 9:案 44-07 逐幀 OSD 時鐘稽核——f805 為 07:25:36,f811 直接跳 07:26:36(剪輯點)')
fig(f'{A}/spot4407.jpg', 13.0,
    '圖 10:案 44-07 標注影片同幀互證——HUD 讀值 55.8 km/h,畫面 OSD 燒錄車速 055Km/h')

h2('4.5 他車測速(申請書「以車追車」情境)')
body('dc007 台86 危險駕駛案為代表性成果:行車紀錄器主人為與事故無關之路人視角,'
     '目標紅色轎車全程 11 秒被追蹤——自匝道加速接近(約 100 km/h)、貼近自車右側'
     ' 10m(約 76 km/h)、貼身過近時接地點被引擎蓋遮擋而誠實無量測、切入前方 30m'
     ' 再加速逃逸;自車速對 GPS 巡航段 MAE ≈ 2.2 km/h。dc006(國道)60 秒內 77 台'
     '目標車獲得絕對速度,分布 62~113 km/h 符合車流常態(圖 11)。')
fig(f'{A}/dc007_t5.jpg', 13.0,
    '圖 11:dc007 危險駕駛案——紅色目標車 13.1m、約 65 km/h;自車視覺碼表 82.8 km/h')

h2('4.6 公開資料集驗證:comma2k19')
body('為使精度可被外部檢驗,本階段於 comma2k19(Schafer 等, 2018;加州灣區通勤,20 Hz,'
     '1164×874,同時提供 CAN 車速與 GNSS/INS/vision 全域位姿)上執行管線 B 自車速評估。'
     '自 Chunk_1 之 188 段中取中位速較高者 24 段 × 60 秒 = 24 分鐘 / 28,703 幀,'
     '取樣點全部由工具自動挑選,無任何人工瞄點。')
body('第一步是反推週期而非套用規範值。以真值反推得 14.37 m(對 CAN)與 14.52 m(對位姿),'
     '與 Caltrans 高速公路標線規格(12 ft 線 + 36 ft 間隔 = 48 ft = 14.63 m)相差 0.7~1.8%,'
     '而與全國 MUTCD 預設值(40 ft = 12.19 m)相差 18%——量測本身即把兩個候選規格分開。'
     '此結果第二次驗證「不可直接套用規範值」的作法(台灣國道規範為 12 m,實測為 10 m)。')
body('以規格值 14.63 m 固定為週期、全程不使用任何真值,結果如表 3。適用條件由 GPS 位置'
     '(路類)與實測路面照度定義,兩者皆不涉及真值。可引用之數字為「高速公路,日夜皆含」'
     '之 MAE 2.79 km/h(覆蓋 87.6%、±3 km/h 內 87.9%);若再限日間則為 1.78 km/h、'
     '中位相對誤差 0.98%,與國內最佳一案(hs005 對人工畫格法 0.69)同一量級,'
     '而對照對象為外部資料集之儀器真值。表 3 亦列出 3.6 節之基頻選擇改動前的數字,'
     '該改動使夜間 MAE 自 12.46 降至 4.72、誤差 p90 自 76.7 降至 3.6,且日間幾乎不受影響。')
add_table(
    ['適用範圍', '覆蓋', '幀數', 'MAE(km/h)', '95% CI', 'p50', 'p90', '±3 內'],
    [
        ['全部 24 段(含市區幹道)', '87.1%', '24,820', '7.73(前 10.54)', '7.48–8.00', '1.15', '7.32', '81.4%'],
        ['高速公路(日夜皆含)', '87.6%', '22,509', '2.79(前 5.32)', '2.65–2.93', '1.05', '3.37', '87.9%'],
        ['高速公路,僅日間', '86.3%', '14,729', '1.78(前 1.84)', '1.71–1.85', '1.03', '3.27', '88.5%'],
        ['高速公路,僅夜間', '90.2%', '7,780', '4.72(前 12.46)', '4.37–5.04', '1.10', '3.58', '86.8%'],
    ],
    [3.9, 1.4, 1.7, 2.6, 2.0, 1.0, 1.0, 1.3],
    caption='表 3:comma2k19 自車速準確度(週期固定 14.63 m、全程不用真值;對 GNSS/INS 位姿真值)')
fig(f'{A}/fig_odo_fix.png', 15.5,
    '圖 12:基頻選擇改為跨時間聯合決策的前後對照——夜間 MAE 12.46 → 4.72、p90 76.7 → 3.6')
body('同一批資料另有一項附帶結論:該車 CAN 車速與 GNSS/INS 位姿之關係為'
     ' pose = 1.0101 × CAN,即 CAN 系統性低報約 1.01%,且為乘性、非雜訊。'
     '本系統估計值對位姿真值之 bias(+1.65)小於對 CAN 之 bias(+2.68),'
     '亦即比該車自身的碼表更接近 GNSS/INS 解。此與 4.2 節「車上儀器自報速度'
     '不是金標準」之結論一致。')

h2('4.7 研究誠信:一次公開更正的誤差歸因')
body('comma2k19 首次評估(2026-08-10)所報端到端 MAE 為 10.54 km/h,並將誤差歸因於'
     '「加州高速公路在兩條漆線之間鋪設反光路釘,造成半週期誤鎖」,且稱該誤鎖受速度閘控。'
     '2026-08-17 重新檢查後確認該歸因的兩個部分都是錯的,本報告予以更正並保留紀錄,'
     '因為此類錯誤的性質本身具方法學價值。')
body('真正的原因在段落挑選:篩選條件為「中位速 ≥ 70 km/h」,而速度並不能代理路類——'
     '該條件放進了一批「起點在有紅綠燈的市區幹道、後段才進入 I-280」的段落,'
     '而市區幹道的法定虛線週期本來就約為高速公路的一半(實測 7.39 m 對 14.63 m)。'
     '將全部段落套用同一常數,跨界段的速度即整段加倍。沿共用走廊依緯度分箱可見(圖 13),'
     '半週期誤鎖率自 0.0%(南段)階梯式上升至 90.4%(北端 1 公里內),'
     '且該處量到的週期正好由 14.5 m 掉到 7.39 m——碼表量到的是路面真實週期,'
     '錯的是我們餵入的常數。原先所稱「速度閘控」則是混淆變項:'
     '幹道段與夜間段剛好也是速度較慢的段落。')
fig(f'{A}/fig_position_step.png', 15.5,
    '圖 13:半週期誤鎖是「位置」的階梯函數而非速度的函數;同一位置量到的週期正好減半')
fig(f'{A}/fig_roadclass.jpg', 15.0,
    '圖 14:同一趟行程的兩處畫面——左為北端市區幹道(紅綠燈、短虛線),右為往南 1 公里的 I-280 本線')
body('此更正對本計畫有兩點意義。第一,誤差歸因錯誤不會被總體數字抓出來:首次評估的'
     '誤差中位數僅 1.20 km/h,表面上相當好,問題只在分群檢查時才顯現;因此'
     '「把誤差分布拆開看」應列為標準程序,而非可選步驟。第二,由此得到一條可事先'
     '計算的適用條件——週期是「每條路的常數」而非「每支影片的常數」,'
     '跨越路類的影片必須事先依路類切分(表 4)。')

h2('4.8 失效邊界(誠實標示)')
body('以下情境系統設計為自動拒發或降級顯示,而非硬給數字:近地平線遠場'
     '(>0.3 m/px 拒發)、螢幕翻拍摩爾紋、貼地視角虛線 3D 取樣、加減速段'
     '(碼表解析度約 3 秒)、換道期間(線位閘拒發)、貼身目標車(接地點低於'
     '引擎蓋線)、LK 光流測速(任何裝置皆不信任)、超短片之片尾窗截斷。'
     '每一項皆以旗標同步呈現於影片與 CSV(far-range、no-lock、灰色問號),'
     '法庭證據寧缺勿錯。本階段另將三條邊界寫成可事先驗算的條件(表 4)。')
add_table(
    ['邊界', '條件', '如何發現', '能否修正'],
    [
        ['取樣率下限', 'fps ≳ 8 × 速度 ÷ 週期(每週期約需 8 幀以上)',
         'KITTI 10 Hz、85 km/h、週期 11.7 m → 僅 5.0 幀/週期,八度選擇崩潰;'
         '該資料集整體 MAE 24.7,但僅計鎖對八度之幀為 2.27,可證壞的是選擇而非量測鏈',
         '不可(取樣率不足)'],
        ['週期是每條路的常數', '影片若跨越路類,須事先依路類切分',
         'comma2k19 市區幹道段量到 7.39 m,套用 14.63 m 即速度加倍(見 4.7)',
         '不可(常數本身錯)'],
        ['夜間逆反射標記', '暗處若存在「漆線週期整數分之一」之次級週期,單一時間窗內無從分辨',
         '同一條 I-280:日間半週期誤鎖率 1.0%、夜間 11.9%',
         '大部分已修(見 3.6;殘留 3.2%)'],
    ],
    [2.6, 3.8, 6.2, 2.9],
    caption='表 4:三條可計算的適用條件')
body('前兩條並非「方法有時候會壞」,而是可事先計算、事先排除的使用條件;'
     '本計畫認為將其明確化本身即為貢獻。國內既有案例全部落在 8.4~15.4 幀/週期,'
     '且每支影片皆為單一路類,未觸及前兩條邊界——這是素材條件使然,'
     '而非設計上的保證,報告中不應以「未曾發生」代替「已被排除」。')

h2('4.9 既有結果的迴歸驗證')
body('3.6 節之基頻選擇修改屬量測邏輯變更,因此對國內 15 個案例(五支示範影片、'
     '20251029 汽車批量六案、海盛與公開影片四案)執行前後逐幀比對。結果:程式重構本身'
     '完全惰性——關閉新選項時,9,762 個共同鎖定幀與改動前程式逐幀完全相同;'
     '開啟後 15 案中 12 案完全零變動,總計 0.74% 的幀有變動,而兩支具備人工逐幀真值'
     '之案例皆為零變動。有變動之幀以「離鄰域共識的距離」判斷方向(不使用真值),'
     '國道一號兩案分別為 15 幀中 14 幀改善、1 幀中 1 幀改善。舊行為保留為命令列開關,'
     '封存輸出隨時可完整重現。')

# ================= 五、結論 =================
h1('五、結論與建議')
body('本計畫完成申請書所列三項目標:(1) 建立了以法定物理錨點為真值的標準化驗證'
     '流程,並以 16 件以上真實案件完成交叉驗證,量化定義了失效邊界;(2) 落實'
     '「AI 幾何 × 物理約束」之尺度校正——形狀信任單幀 3D 重建,尺度信任台灣法定'
     '標線,校正變異以 CV 傳遞至速度信賴區間;(3) 在移動式行車紀錄器上以虛線週期'
     '碼表解決自車速,以「自車速 + 距離變化率」重建他車絕對速度,並於危險駕駛'
     '真實影片完成全程重建示範。')
body('主要量化成果:CCTV 對人工交比法 +0.5%;行車紀錄器自車速對人工畫格法 MAE'
     ' 0.69~1.30 km/h、對 GPS 為 1~2.4 km/h;8 件鑑定案 28 段批量對人工畫格法'
     ' MAE 1.93(可信 23 段 1.37)km/h;公開資料集 comma2k19 之 22,509 幀對'
     ' GNSS/INS 位姿真值 MAE 2.79 km/h(僅日間 1.78,中位相對誤差 0.98%)。'
     '此外,「證物影片時基三陷阱」之發現與稽核流程,對現行人工鑑定流程亦有直接'
     '改進價值——建議任何速度計算前,一律以畫面內 OSD 時鐘或 GPS 重建時間軸。')
body('後續建議:(1) 管線 A 仍缺公開資料集驗證。BrnoCompSpeed 為文獻之標準比較對象'
     '且附相機標定(可將標定誤差與量測鏈誤差分離),惟資料量約 200 GB 且其官方主機'
     '於本研究環境無法連線,建議透過校內網路或由指導教授具名向作者索取子集;'
     '(2) comma2k19 目前僅使用一個 chunk,即單一車輛與單一走廊,擴充至其餘 chunk'
     '可取得多路網驗證;(3) 若要微調偵測器,正確的資料集是具 instance mask 的 nuImages'
     '而非僅有框者,因本系統之接地點取自遮罩最低 12% 而非框底邊,'
     '僅微調偵測頭對量測精度助益有限;(4) 取得更多「汽車目標 + 路人視角」之事故影片'
     '以擴充他車測速之實案驗證;(5) 整理程式碼並發表會議或期刊論文。')

# ================= 參考文獻 =================
h1('參考文獻')
refs = [
    '[1] 道路交通安全督導委員會,道安資訊查詢網事故統計。',
    '[2] 交通部運輸資料流通服務平臺(TDX),即時路口 CCTV 影像。',
    '[3] Li, Z., & Snavely, N. (2018). MegaDepth: Learning Single-View Depth Prediction from Internet Photos. CVPR.',
    '[4] Hu, M., et al. (2024). Metric3D v2: A Versatile Monocular Geometric Foundation Model. TPAMI.',
    '[5] Yang, L., et al. (2024). Depth Anything V2. NeurIPS.',
    '[6] Piccinelli, L., et al. (2025). UniDepth v2: Universal Monocular Metric Depth Estimation. arXiv.',
    '[7] Apple Machine Learning Research (2025). SHARP: 單幀 3D 高斯重建開源模型(ml-sharp)。',
    '[8] Ultralytics (2023). YOLOv8: Real-Time Object Detection and Segmentation.',
    '[9] 內政部與交通部,道路交通標誌標線號誌設置規則(車道線、行車分向線與導引線規格)。',
    '[10] Sochor, J., et al. (2019). Comprehensive Data Set for Automatic Single Camera Visual '
    'Speed Measurement. IEEE Transactions on Intelligent Transportation Systems, 20(5).',
    '[11] Luvizon, D. C., Nassu, B. T., & Minetto, R. (2017). A Video-Based System for Vehicle '
    'Speed Measurement in Urban Roadways. IEEE Transactions on Intelligent Transportation Systems.',
    '[12] Schafer, H., Santana, E., Haden, A., & Biasini, R. (2018). A Commute in Data: '
    'The comma2k19 Dataset. arXiv:1812.05752.',
    '[13] Geiger, A., Lenz, P., & Urtasun, R. (2012). Are we ready for Autonomous Driving? '
    'The KITTI Vision Benchmark Suite. CVPR.',
    '[14] Caesar, H., et al. (2020). nuScenes: A Multimodal Dataset for Autonomous Driving. CVPR.',
    '[15] Fernández Llorca, D., Hernández Martínez, A., & García Daza, I. (2021). '
    'Vision-based vehicle speed estimation: A survey. IET Intelligent Transport Systems, 15(8). '
    '(arXiv:2101.06159)',
    '[16] California Department of Transportation. Standard Plan A20A: Pavement Markings '
    '(Lane Line, Freeway Application).',
    '[17] Federal Highway Administration. Manual on Uniform Traffic Control Devices (MUTCD), '
    'Part 3: Markings.',
]
for r in refs:
    para(r, size=11, space_after=3, line=1.3)

out = f'{OUT_DIR}/成果報告草稿_20260817.docx'
doc.save(out)
print('saved:', out)
