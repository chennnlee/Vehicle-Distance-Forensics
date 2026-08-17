"""Structural check on a generated deck, standing in for looking at it.

This box has no LibreOffice and no CJK font, so the deck cannot be rendered here
and eyeballed. Three classes of defect are still findable from the file itself,
and they are the ones hand-written slide geometry actually produces:

  bounds      a shape running off the 13.33 x 7.5 in canvas
  distortion  an image whose placed aspect ratio differs from the file's, which
              is invisible in code and obvious on screen
  overflow    a text box holding more characters than its area can show at the
              font size it asks for

The overflow estimate is deliberately crude -- a CJK glyph is about one em wide,
so characters-per-line is roughly box_width / font_size, and lines are about
1.5 em tall. It flags candidates to re-read, not a verdict.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

EMU_PER_IN = 914400


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--fill", type=float, default=0.92,
                    help="Flag a text box when the estimated fill exceeds this share of its height.")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    W = prs.slide_width / EMU_PER_IN
    H = prs.slide_height / EMU_PER_IN
    print(f"{Path(args.pptx).name}:  {len(prs.slides)} 張,版面 {W:.2f} x {H:.2f} in\n")
    problems = 0

    for i, slide in enumerate(prs.slides, 1):
        issues = []
        n_img = n_txt = n_tbl = 0
        for sh in slide.shapes:
            x, y = sh.left / EMU_PER_IN, sh.top / EMU_PER_IN
            w = (sh.width or 0) / EMU_PER_IN
            h = (sh.height or 0) / EMU_PER_IN
            if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
                issues.append(f"越界 {sh.shape_type} 在 ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}")
            if sh.shape_type == 13 or getattr(sh, "image", None) is not None:  # PICTURE
                n_img += 1
                try:
                    blob = sh.image.blob
                    from io import BytesIO
                    iw, ih = Image.open(BytesIO(blob)).size
                    want, got = iw / ih, w / h
                    if abs(want - got) / want > 0.02:
                        issues.append(f"圖片變形：檔案比例 {want:.3f} vs 放置比例 {got:.3f} "
                                      f"（差 {100*abs(want-got)/want:.1f}%）→ 建議 h={w/want:.2f}")
                except Exception as exc:                     # noqa: BLE001
                    issues.append(f"圖片無法讀取比例：{exc}")
            if sh.has_text_frame and sh.text_frame.text.strip():
                n_txt += 1
                chars = sum(len(r.text) for pa in sh.text_frame.paragraphs for r in pa.runs)
                sizes = [r.font.size.pt for pa in sh.text_frame.paragraphs for r in pa.runs
                         if r.font.size is not None]
                fs = max(sizes) if sizes else 18.0
                per_line = max(1.0, (w * 72.0) / fs)          # CJK glyph ~= 1 em wide
                n_para = len([pa for pa in sh.text_frame.paragraphs if pa.text.strip()])
                lines = sum(max(1, round(len(pa.text) / per_line + 0.4))
                            for pa in sh.text_frame.paragraphs if pa.text.strip())
                need_in = lines * fs * 1.5 / 72.0
                if h > 0 and need_in > args.fill * h:
                    issues.append(f"文字可能溢出：{chars} 字 / {n_para} 段 → 約 {lines} 行 "
                                  f"需 {need_in:.2f} in,框高 {h:.2f} in")
            if sh.has_table:
                n_tbl += 1
        flag = "⚠" if issues else "✓"
        print(f"{flag} 第 {i:2d} 張  文字框 {n_txt} 圖 {n_img} 表 {n_tbl}")
        for it in issues:
            print(f"      - {it}")
            problems += 1
    print(f"\n可疑項目共 {problems} 個")


if __name__ == "__main__":
    main()
